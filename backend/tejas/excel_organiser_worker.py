"""Bounded, text-only XLSX parsing. No formula evaluation, macros or external links."""
import codecs
import io
import json
import os
import re
import sys
import zipfile
import openpyxl
from pdf_organiser_worker import seal

MAX_FILE = 5 * 1024 * 1024
MAX_ROWS = 300
HEADERS = {'question_id', 'question_no', 'subject', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option', 'solution_text', 'topic', 'difficulty', 'type', 'exam', 'duration_minutes', 'marks', 'negative_marks', 'source'}


def workbook(raw):
    if not raw or len(raw) > MAX_FILE:
        raise ValueError('Use an .xlsx workbook of at most 5 MB.')
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        entries = z.infolist()
        if len(entries) > 200 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
            raise ValueError('Workbook expansion exceeds 32 MB. Split the workbook.')
        names = [e.filename.lower() for e in entries]
        if any(any(t in n for t in ('vbaproject', 'externallink', 'xl/media/', 'xl/embeddings/')) for n in names):
            raise ValueError('Use a text-only workbook without macros, external links or embedded images.')
        if len(names) != len(set(names)):
            raise ValueError('Workbook contains duplicate entries.')
    return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)


def organise(qraw, sraw=None):
    issues = []
    def issue(sheet, row, field, message):
        issues.append({'sheet': sheet, 'row': row, 'field': field, 'message': message, 'question': None})
    def rows(ws, role):
        if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > 24:
            raise ValueError('Use at most 300 data rows and 24 columns per sheet.')
        it = ws.iter_rows()
        first = next(it, [])
        headers = [str(c.value or '').strip().lower() for c in first]
        while headers and not headers[-1]: headers.pop()
        if not headers:
            issue(ws.title, 1, 'headers', 'Add the template headers.'); return []
        if len(headers) != len(set(headers)) or any(h not in HEADERS for h in headers):
            issue(ws.title, 1, 'headers', 'Unknown or duplicate columns. Use the downloadable template.'); return []
        required = {'question_no', 'question_text'} if role == 'questions' else {'question_no', 'solution_text'}
        for h in required - set(headers): issue(ws.title, 1, h, 'Required column is missing.')
        result = []
        for number, cells in enumerate(it, 2):
            if number > MAX_ROWS + 1: raise ValueError('Use at most 300 data rows per sheet.')
            if not any(c.value is not None for c in cells): continue
            r = {'_row': number, '_sheet': ws.title}
            for h, c in zip(headers, cells):
                if c.data_type in ('f', 'e'):
                    issue(ws.title, number, h, 'Formulas and error cells are not imported. Paste values instead.')
                value = '' if c.value is None else str(c.value).strip()
                if len(value) > 12000: issue(ws.title, number, h, 'Cell exceeds 12,000 characters.')
                r[h] = value[:12000]
            no = r.get('question_no', '')
            if not re.fullmatch(r'[1-9]\d{0,4}', no):
                issue(ws.title, number, 'question_no', 'Use a positive whole question number.')
            result.append(r)
        return result
    qbook = workbook(qraw)
    sbook = workbook(sraw) if sraw else None
    allowed = {'Questions', 'Solutions', 'Instructions'}
    for book in [qbook] + ([sbook] if sbook else []):
        if len(book.sheetnames) > 3 or any(n not in allowed for n in book.sheetnames):
            raise ValueError('Name sheets Questions and Solutions (optional Instructions). Use the template.')
    if 'Questions' not in qbook.sheetnames: raise ValueError('The first workbook needs a Questions sheet.')
    qs = rows(qbook['Questions'], 'questions')
    source = sbook or qbook
    ss = rows(source['Solutions'], 'solutions') if 'Solutions' in source.sheetnames else None
    if sbook and ss is None: raise ValueError('The second workbook needs a Solutions sheet.')
    if not qs: issue('Questions', 2, 'question_text', 'Add at least one question.')
    for group in [qs] + ([ss] if ss is not None else []):
        for field in ['question_id', 'question_no']:
            seen = set()
            for r in group:
                v = r.get(field)
                if v and v in seen: issue(r['_sheet'], r['_row'], field, 'Duplicate identity; each question and solution must be unique.')
                if v: seen.add(v)
    output = []; used = set(); exams = set(); durations = set(); marks = set(); negatives = set()
    for order, q in enumerate(qs, 1):
        before = len(issues); s = q
        if ss is not None:
            by_id = [r for r in ss if q.get('question_id') and r.get('question_id') == q['question_id']]
            by_no = [r for r in ss if r.get('question_no') == q.get('question_no')]
            candidates = by_id or by_no
            s = candidates[0] if len(candidates) == 1 else {}
            if not s:
                issue('Questions', q['_row'], 'solution_text', 'No unique matching solution.')
            elif (s.get('question_no') != q.get('question_no') or (q.get('question_id') and s.get('question_id') and q['question_id'] != s['question_id'])):
                issue('Questions', q['_row'], 'question_id', 'Question ID and number disagree with the solution.'); s = {}
            else: used.add(s['_row'])
        kind = q.get('type') or 'mcq'
        if kind not in ('mcq', 'integer'): issue('Questions', q['_row'], 'type', 'Use mcq or integer.')
        answer = q.get('correct_option') or s.get('correct_option', '')
        if q.get('correct_option') and s.get('correct_option') and q['correct_option'] != s['correct_option']:
            issue('Questions', q['_row'], 'correct_option', 'Question and solution answer keys conflict.')
        options = {k: q.get('option_' + k.lower(), '') for k in 'ABCD'} if kind == 'mcq' else {}
        if kind == 'mcq' and (answer not in 'ABCD' or len(answer) != 1 or not all(options.values())):
            issue('Questions', q['_row'], 'correct_option', 'MCQs require four options and an explicit A, B, C or D key.')
        if kind == 'integer':
            if not re.fullmatch(r'[+-]?\d{1,20}', answer): issue('Questions', q['_row'], 'correct_option', 'Enter an exact integer key (no decimals).')
            else: answer = str(int(answer))
        if not q.get('question_text'): issue('Questions', q['_row'], 'question_text', 'Question text is required.')
        if not s.get('solution_text'): issue(s.get('_sheet', 'Questions'), s.get('_row', q['_row']), 'solution_text', 'Matched solution text is required.')
        if ss is not None and q.get('solution_text') and q['solution_text'] != s.get('solution_text'):
            issue('Questions', q['_row'], 'solution_text', 'Combined and separate solutions conflict. Keep one source.')
        exams.add(q.get('exam') or 'JEE Main'); durations.add(q.get('duration_minutes') or '180')
        if q.get('marks'): marks.add(q['marks'])
        if q.get('negative_marks'): negatives.add(q['negative_marks'])
        output.append({'number': order, 'source_number': q.get('question_no'), 'source_id': q.get('question_id', ''),
            'solution_number': order if s.get('solution_text') else None, 'subject': q.get('subject') or 'General',
            'section': q.get('subject') or 'General', 'type': kind, 'text': q.get('question_text', ''), 'options': options,
            'answer': answer, 'solution': s.get('solution_text', ''), 'topic': q.get('topic', ''), 'difficulty': q.get('difficulty', ''),
            'render_mode': 'typed', 'q_refs': [], 's_refs': [], 'issues': [x['message'] for x in issues[before:]],
            'pairing_evidence': f"Excel Questions row {q['_row']} → {s.get('_sheet', 'Questions')} row {s.get('_row', q['_row'])}",
            'source_row': q['_row']})
    if ss is not None:
        for s in ss:
            if s['_row'] not in used: issue('Solutions', s['_row'], 'question_no', 'Unmatched solution row.')
    exam = next(iter(exams), 'JEE Main')
    if len(exams) != 1 or exam not in ('JEE Main', 'BITSAT'): issue('Questions', 1, 'exam', 'Use one exam per workbook: JEE Main or BITSAT.')
    if exam == 'BITSAT' and any(q['type'] != 'mcq' for q in output): issue('Questions', 1, 'type', 'BITSAT supports single-correct MCQs only.')
    duration = next(iter(durations), '180')
    if len(durations) != 1 or not duration.isdigit() or not 1 <= int(duration) <= 360:
        issue('Questions', 1, 'duration_minutes', 'Use one duration of 1–360 minutes for the test.'); duration = '180'
    correct = 3 if exam == 'BITSAT' else 4
    if marks and marks != {str(correct)}: issue('Questions', 1, 'marks', f'This practice importer uses +{correct} marks per correct answer.')
    if negatives and negatives != {'-1'}: issue('Questions', 1, 'negative_marks', 'This practice importer uses -1 per incorrect answer.')
    result = {'questions': output, 'issues': issues, 'exam': exam, 'duration_minutes': int(duration),
        'source_pages': {}, 'sandbox': 'chroot_low_uid_seccomp_resource_limits'}
    if len(json.dumps(result).encode()) > 6 * 1024 * 1024: raise ValueError('Imported text exceeds 6 MB. Split the test.')
    return result


def main():
    try:
        directory = sys.argv[1]
        with open(os.path.join(directory, 'questions.xlsx'), 'rb') as f: qraw = f.read(MAX_FILE + 1)
        path = os.path.join(directory, 'solutions.xlsx')
        sraw = None
        if os.path.exists(path):
            with open(path, 'rb') as f: sraw = f.read(MAX_FILE + 1)
        # Warm trusted XML/ZIP paths before sealing filesystem access.
        for name in ('utf-8', 'utf-16', 'cp437', 'raw_unicode_escape'): codecs.lookup(name)
        warm = io.BytesIO(); b = openpyxl.Workbook(); b.save(warm)
        w = openpyxl.load_workbook(io.BytesIO(warm.getvalue()), read_only=True)
        list(w.active.iter_rows()); w.close()
        seal(directory)
        result = organise(qraw, sraw)
        sys.stdout.write(json.dumps(result, ensure_ascii=False))
    except ValueError as exc:
        sys.stdout.write(json.dumps({'error': str(exc)[:300]})); sys.exit(2)
    except Exception:
        sys.stdout.write(json.dumps({'error': 'This workbook could not be read safely. Use an unlocked .xlsx file exported from the template.'})); sys.exit(3)

if __name__ == '__main__': main()
