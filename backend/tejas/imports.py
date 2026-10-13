"""Unified importer. Immutable sources, fenced one-shot processing, single-document publication.
The pilot uses request-triggered background dispatch, not a distributed worker fleet.
Interrupted processing is recoverable by an explicit, lease-fenced retry.
"""
import asyncio
import base64
import csv
import hashlib
import io
import json
from pathlib import Path
import tempfile
import uuid
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, Request, BackgroundTasks, Response
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError
from starlette.concurrency import run_in_threadpool
from .auth import owner, staff
from .db import db, sets, files, uid, now, get, clean
from .import_models import ImportView, ImportList, ScheduleInput, LiveInput
from .sandbox import organise_pdfs
from .content import valid_rules, question_issues
from tejas_storage import put_object, get_object

router = APIRouter(prefix='/api/studio/imports', tags=['Unified test importer'])
imports = db.tejas_imports
SLOTS = asyncio.Semaphore(2)
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def view(record, s=None, detail=True):
    s = s or {}
    status = record['status']
    if status == 'organising' and record.get('lease_until', 0) < now(): status = 'interrupted'
    if s.get('start_at'):
        status = 'ended' if now() >= s['end_at'] else 'unlive' if not s.get('live_enabled') else 'scheduled' if now() < s['start_at'] else 'live'
    return ImportView(id=record['id'], title=record['title'], format=record['format'], status=status,
        message=record.get('message', ''), questions=record.get('questions', []) if detail else [],
        issues=record.get('issues', []) if detail else [], count=len(record.get('questions', [])),
        exam=record.get('exam', 'JEE Main'), duration_minutes=record.get('duration_minutes', 180),
        set_id=s.get('id'), start_at=s.get('start_at'), end_at=s.get('end_at'),
        live_enabled=s.get('live_enabled', False), revision=s.get('revision', 0), server_now=now())


async def read_record(id):
    record = await imports.find_one({'id': id}, {'_id': 0})
    if not record: raise HTTPException(404, 'Import not found.')
    return record


@router.get('/template')
async def template(variant: str = 'combined', u=Depends(staff)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    if variant not in ('combined', 'sheets', 'questions', 'solutions'): raise HTTPException(400, 'Unknown template layout.')
    b = Workbook(); b.remove(b.active)
    qh = ['question_id','question_no','subject','question_text','option_a','option_b','option_c','option_d','correct_option','type','topic','difficulty','exam','duration_minutes']
    qr = ['phy-1',1,'Physics','What is the SI unit of force?','Newton','Joule','Watt','Pascal','A','mcq','Units','Easy','JEE Main',180]
    if variant != 'solutions':
        w = b.create_sheet('Questions'); w.append(qh + (['solution_text'] if variant=='combined' else [])); w.append(qr + (['Force is measured in newtons (N).'] if variant=='combined' else []))
    if variant in ('sheets','solutions'):
        w=b.create_sheet('Solutions'); w.append(['question_id','question_no','solution_text','correct_option']); w.append(['phy-1',1,'Force is measured in newtons (N).','A'])
    w=b.create_sheet('Instructions')
    for line in ['Replace the example row; maximum 300 questions and 5 MB per workbook.', 'Keep sheet names and headers unchanged. Plain values only: no formulas, macros or embedded images.', 'Question IDs and numbers must be unique and agree across sheets/files. Solutions and explicit keys are required.', 'Use mcq with A–D options and keys, or integer with an exact integer key in correct_option.', 'Optional exam: JEE Main or BITSAT (MCQ only). Default duration: 180 minutes. One exam and duration per workbook.', 'Custom practice rules: JEE Main +4 / -1; BITSAT +3 / -1. Not an official-year simulation.', 'Use $LaTeX$ for basic equations. Unsupported report notation is flagged; nothing is guessed.', 'Start date/time is chosen after import. Templates contain example content; remove it before importing your own paper.']:
        w.append([line])
    for w in b:
        w.freeze_panes='A2'
        for c in w[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='5C0025')
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width=36
            for c in col: c.alignment=Alignment(wrap_text=True, vertical='top')
    out=io.BytesIO(); b.save(out)
    return Response(out.getvalue(),media_type=XLSX,headers={'Content-Disposition':f'attachment; filename="Tejas_{variant}_template.xlsx"'})


@router.get('', response_model=ImportList)
async def list_imports(u=Depends(staff)):
    records=await imports.find({'qa_only': {'$ne': True}}, {'_id':0}).sort('created_at',-1).limit(50).to_list(50)
    linked=await sets.find({'id':{'$in':[r['id'] for r in records]}}, {'_id':0,'id':1,'start_at':1,'end_at':1,'revision':1,'live_enabled':1}).to_list(50)
    by_id={s['id']:s for s in linked}
    return ImportList(items=[view(r,by_id.get(r['id']),False) for r in records])


@router.post('', response_model=ImportView, status_code=202)
async def create_import(request: Request, bg: BackgroundTasks, u=Depends(owner)):
    # Middleware authenticates and bounds multipart bytes before FastAPI can spool.
    async with request.form(max_files=2,max_fields=2) as form:
        if set(form.keys()) - {'questions','solutions','request_id','format'} or len(form.multi_items()) != len(form.keys()):
            raise HTTPException(422,'Use one questions file and, optionally, one solutions file.')
        format=str(form.get('format',''))
        if format not in ('pdf','xlsx'): raise HTTPException(422,'Choose PDF or Excel (.xlsx).')
        try: request_id=str(uuid.UUID(str(form.get('request_id',''))))
        except ValueError: raise HTTPException(422,'A UUID upload identifier is required.')
        id=str(uuid.uuid5(uuid.NAMESPACE_URL,u['id']+':import:'+request_id))
        sources=[]; payloads=[]
        for role in ['questions','solutions']:
            f=form.get(role)
            if not f:
                if role=='questions' or format=='pdf': raise HTTPException(422,'Upload both Questions and Solutions PDFs, or an Excel workbook.')
                continue
            if not hasattr(f,'read') or not (f.filename or '').lower().endswith('.'+format): raise HTTPException(422,'File extension must match the selected format.')
            limit=(50 if format=='pdf' else 5)*1024*1024
            raw=await f.read(limit+1)
            if not raw or len(raw)>limit: raise HTTPException(413,f'Use {format.upper()} files within the displayed size limit.')
            if not raw.startswith(b'%PDF-' if format=='pdf' else b'PK'): raise HTTPException(422,'File content does not match its format.')
            sources.append({'id':uid(),'kind':role,'name':Path(f.filename).name[:160], 'sha256':hashlib.sha256(raw).hexdigest(),'size':len(raw)})
            payloads.append(raw)
        fingerprint=hashlib.sha256(json.dumps([(s['kind'],s['sha256']) for s in sources]).encode()+format.encode()).hexdigest()
        existing=await imports.find_one({'id':id},{'_id':0})
        if existing:
            if existing['fingerprint']!=fingerprint: raise HTTPException(409,'Upload identifier already belongs to different files. Start a new upload.')
            return view(existing,await sets.find_one({'id':id},{'_id':0}))
        if SLOTS.locked(): raise HTTPException(429,'Two imports are processing. Please retry shortly.',headers={'Retry-After':'5'})
        async with SLOTS:
            for source,raw in zip(sources,payloads):
                source['path']=await run_in_threadpool(put_object,f"tejas/imports/{id}/{source['id']}.{format}",raw,'application/pdf' if format=='pdf' else XLSX)
        title=Path(sources[0]['name']).stem.replace('_',' ').strip()[:150] or 'Imported practice test'
        record={'id':id,'title':title,'format':format,'fingerprint':fingerprint,'sources':sources,'status':'queued',
                'created_at':now(),'created_by':u['id'],'questions':[],'issues':[],
                'rights_attested_at':now(),'rights_attested_by':u['id']}
        # Atomic identity claim; duplicate POST cannot create another import/test.
        await imports.update_one({'id':id},{'$setOnInsert':record},upsert=True)
        bg.add_task(process_import,id)
        return view(record)


async def process_import(id):
    async with SLOTS:
        token=uid()
        r=await imports.find_one_and_update({'id':id,'$or':[{'status':{'$in':['queued','failed']}},{'status':'organising','lease_until':{'$lt':now()}}]},
            {'$set':{'status':'organising','lease_token':token,'lease_until':now()+900,'message':'Reading sources, pairing questions, and checking explicit answer keys…'}},return_document=ReturnDocument.AFTER)
        if not r: return
        fence={'id':id,'lease_token':token,'status':'organising'}
        try:
            with tempfile.TemporaryDirectory(prefix='tejas-import-') as directory:
                paths={}
                for s in r['sources']:
                    raw=await run_in_threadpool(get_object,s['path'])
                    if hashlib.sha256(raw).hexdigest()!=s['sha256']: raise ValueError('Source integrity check failed. Upload again.')
                    p=Path(directory)/(s['kind']+'.'+r['format']); p.write_bytes(raw); paths[s['kind']]=p
                result=await organise_pdfs(paths['questions'],paths.get('solutions'),r['format'])
            qs=result['questions']; issues=result['issues']
            for q in qs:
                q['id']=str(uuid.uuid5(uuid.UUID(id),str(q['number'])))
                q['approved']=False  # Structural validation is not academic reviewer approval.
                q['accessible_text']=q['text']; q['solution_accessible_text']=q['solution']
                if r['format']=='pdf':
                    assets=q.pop('assets',{})
                    for role,key in [('question','q_refs'),('solution','s_refs')]:
                        original=next(s for s in r['sources'] if s['kind']==('questions' if role=='question' else 'solutions'))
                        refs=q[key]
                        for n,encoded in enumerate(assets.get(role,[])):
                            png=base64.b64decode(encoded,validate=True)
                            fid=str(uuid.uuid5(uuid.UUID(q['id']),role+':'+str(n)))
                            path=await run_in_threadpool(put_object,f'tejas/imports/{id}/renders/{token}/{fid}.png',png,'image/png')
                            await files.update_one({'id':fid},{'$setOnInsert':{'id':fid,'set_id':id,'path':path,'kind':'rendered_png','sha256':hashlib.sha256(png).hexdigest(),'source_sha256':original['sha256'],'created_at':now(),'is_deleted':False}},upsert=True)
                            refs[n]['file_id']=fid; refs[n]['prerendered']=True
                for message in question_issues(q,result['exam']):
                    issues.append({'question':q['number'],'row':q.get('source_row'),'message':message})
                if q['render_mode']=='typed':
                    from .reports import validate_typed_math
                    try: await run_in_threadpool(validate_typed_math,q)
                    except ValueError as exc: issues.append({'question':q['number'],'row':q.get('source_row'),'field':'math','message':str(exc)})
            if len(json.dumps(qs,ensure_ascii=False).encode())>6*1024*1024: raise ValueError('Test content exceeds 6 MB. Split the source.')
            await imports.update_one(fence,{'$set':{'questions':qs,'issues':issues,'exam':result['exam'],'duration_minutes':result['duration_minutes'],
                'status':'blocked' if issues else 'ready','sandbox':result['sandbox'],
                'message':'Resolve the listed source issues and upload again.' if issues else 'All source entries paired. Supplied answer keys retained; no answers were generated.'}})
        except Exception as exc:
            from tejas_storage import StorageError
            message=str(exc)[:350] if isinstance(exc,(ValueError,StorageError)) else 'Processing interrupted. Retry with the same saved sources.'
            await imports.update_one(fence,{'$set':{'status':'failed','message':message}})


@router.get('/{id}',response_model=ImportView)
async def read_import(id:str,u=Depends(staff)):
    r=await read_record(id)
    return view(r,await sets.find_one({'id':id},{'_id':0}))


@router.post('/{id}/retry',response_model=ImportView)
async def retry_import(id:str,bg:BackgroundTasks,u=Depends(owner)):
    r=await read_record(id)
    if r['status'] not in ('queued','failed') and not (r['status']=='organising' and r.get('lease_until',0)<now()):
        raise HTTPException(409,'This import is already processing or complete. Replace invalid sources with a new upload.')
    bg.add_task(process_import,id)
    return view(r)


@router.get('/{id}/errors')
async def error_report(id:str,u=Depends(staff)):
    r=await read_record(id); stream=io.StringIO(); writer=csv.writer(stream)
    writer.writerow(['Sheet','Row','Question','Field','Issue'])
    for e in r.get('issues',[]):
        values=[e.get(k,'') for k in ('sheet','row','question','field','message')]
        writer.writerow(["'"+str(v) if str(v).startswith(('=','+','-','@')) else v for v in values])
    return Response(stream.getvalue(),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="Tejas_import_issues.csv"'})


@router.get('/{id}/questions/{qid}/asset/{kind}/{index}')
async def import_asset(id:str,qid:str,kind:str,index:int,u=Depends(staff)):
    r=await read_record(id)
    q=next((q for q in r.get('questions',[]) if q['id']==qid),None)
    refs=(q or {}).get('q_refs' if kind=='question' else 's_refs',[])
    if kind not in ('question','solution') or not 0<=index<len(refs): raise HTTPException(404,'Source preview not found.')
    f=await get(files,refs[index]['file_id'])
    return Response(await run_in_threadpool(get_object,f['path']),media_type='image/png')


@router.post('/{id}/schedule',response_model=ImportView)
async def schedule_import(id:str,data:ScheduleInput,u=Depends(owner)):
    r=await read_record(id)
    if r['status']!='ready' or r.get('issues') or not r.get('questions'): raise HTTPException(409,'Resolve all source issues before scheduling.')
    try:
        date=datetime.fromisoformat(data.start_at.replace('Z','+00:00'))
        if date.tzinfo is None: raise ValueError()
        start=date.astimezone(timezone.utc).timestamp()
        if not now()-60<=start<=now()+366*86400: raise ValueError()
    except ValueError: raise HTTPException(422,'Choose a start time from now to one year ahead, including timezone.')
    existing=await sets.find_one({'id':id},{'_id':0})
    end=start+r['duration_minutes']*60
    if existing:
        if existing.get('start_at',0)<=now(): raise HTTPException(409,'A started test cannot be rescheduled. Upload a new test for another competition.')
        res=await sets.update_one({'id':id,'revision':data.revision,'start_at':{'$gt':now()}},{'$set':{'start_at':start,'end_at':end,'live_enabled':True,'status':'published'},'$inc':{'revision':1}})
        if not res.modified_count: raise HTTPException(409,'Schedule changed elsewhere. Reload before retrying.')
    else:
        if data.revision!=0: raise HTTPException(409,'Reload this import before scheduling.')
        rules=valid_rules({'minutes':r['duration_minutes'],'correct':3 if r['exam']=='BITSAT' else 4,'wrong':-1,'skipped':0})
        s={'id':id,'import_id':id,'title':r['title'],'exam':r['exam'],'year':'','source':'Owner-supplied '+r['format'].upper()+'; deterministic ID/number pairing; supplied keys, not independently academically verified.',
           'attribution':'Owner-supplied study material','questions':r['questions'],'file_ids':[ref['file_id'] for q in r['questions'] for key in ('q_refs','s_refs') for ref in q.get(key,[])],
           'rules':rules,'status':'published','version':1,'revision':1,'created_at':r['created_at'],'published_at':now(),
           'profile_mode':'practice','audience':'all','allowed_phones':[],'allowed_emails':[],'expires':'',
           'rights_upload':True,'rights_publish':True,'rights_report':True,'rights_attested_by':r['rights_attested_by'],
           'start_at':start,'end_at':end,'live_enabled':True,'competition':True}
        # One immutable content version in one insert; no archive-before-CAS race.
        try: await sets.insert_one(s)
        except DuplicateKeyError: raise HTTPException(409,'Another schedule was saved. Reload to see it.')
    return view(r,await sets.find_one({'id':id},{'_id':0}))


@router.post('/{id}/live',response_model=ImportView)
async def toggle_live(id:str,data:LiveInput,u=Depends(owner)):
    r=await read_record(id)
    if r['status']!='ready': raise HTTPException(409,'Only a ready test can be enabled.')
    s=await sets.find_one({'id':id},{'_id':0})
    if not s: raise HTTPException(409,'Choose the attempt start date/time first.')
    if data.live and now()>=s['end_at']: raise HTTPException(409,'This test window has ended. Create a new test for a new competition.')
    changed=await sets.update_one({'id':id,'revision':data.revision},{'$set':{'live_enabled':data.live,'status':'published' if data.live else 'withdrawn'},'$inc':{'revision':1}})
    if not changed.modified_count: raise HTTPException(409,'Another change was saved. Reload before retrying.')
    return view(r,await sets.find_one({'id':id},{'_id':0}))
