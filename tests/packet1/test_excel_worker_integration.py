"""Test Excel organiser worker through sandbox (realistic invocation).

Tests combined/two sheets/two files layouts with validation.
"""
import asyncio
import io
import shutil
import tempfile
from pathlib import Path
import pytest
import openpyxl

# Use sandbox to invoke worker
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'backend'))
from tejas.sandbox import organise_pdfs


def create_workbook_bytes(sheets_data):
    """Create Excel workbook bytes from sheet data.
    
    Args:
        sheets_data: Dict of {sheet_name: [headers, row1, row2, ...]}
    
    Returns:
        bytes of the workbook
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_excel_combined_sheet_success():
    """Test successful combined sheet with questions and solutions together."""
    sheets = {
        'Questions': [
            ['question_id', 'question_no', 'subject', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option', 'type', 'solution_text', 'exam', 'duration_minutes'],
            ['phy-1', '1', 'Physics', 'What is F=ma?', 'Newton law', 'Einstein law', 'Bohr law', 'Planck law', 'A', 'mcq', 'F=ma is Newton second law', 'JEE Main', '180'],
            ['chem-1', '2', 'Chemistry', 'What is H2O?', 'Water', 'Oxygen', 'Hydrogen', 'Nitrogen', 'A', 'mcq', 'H2O is water', 'JEE Main', '180'],
        ],
        'Instructions': [['Template instructions']]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result, f"Unexpected error: {result.get('error')}"
        assert result.get('sandbox') == 'chroot_low_uid_seccomp_resource_limits'
        assert len(result['questions']) == 2
        assert len(result['issues']) == 0
        
        # Verify Q1
        q1 = result['questions'][0]
        assert q1['source_number'] == '1'
        assert q1['subject'] == 'Physics'
        assert q1['type'] == 'mcq'
        assert q1['answer'] == 'A'
        assert q1['text'] == 'What is F=ma?'
        assert q1['solution'] == 'F=ma is Newton second law'
        assert len(q1['options']) == 4
        assert q1['options']['A'] == 'Newton law'
        
        # Verify Q2
        q2 = result['questions'][1]
        assert q2['source_number'] == '2'
        assert q2['subject'] == 'Chemistry'
        assert q2['answer'] == 'A'
        
        # Verify exam and duration
        assert result['exam'] == 'JEE Main'
        assert result['duration_minutes'] == 180
        
        print("✅ Excel combined sheet test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_two_sheets_same_workbook():
    """Test Questions and Solutions in separate sheets of same workbook."""
    sheets = {
        'Questions': [
            ['question_id', 'question_no', 'subject', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'type', 'exam', 'duration_minutes'],
            ['math-1', '1', 'Mathematics', 'What is 2+2?', '3', '4', '5', '6', 'mcq', 'JEE Main', '180'],
            ['math-2', '2', 'Mathematics', 'What is 10*5?', '', '', '', '', 'integer', 'JEE Main', '180'],
        ],
        'Solutions': [
            ['question_id', 'question_no', 'solution_text', 'correct_option'],
            ['math-1', '1', '2+2 equals 4', 'B'],
            ['math-2', '2', '10*5 equals 50', '50'],
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert len(result['questions']) == 2
        assert len(result['issues']) == 0
        
        # Verify Q1 MCQ
        q1 = result['questions'][0]
        assert q1['type'] == 'mcq'
        assert q1['answer'] == 'B'
        assert q1['solution'] == '2+2 equals 4'
        
        # Verify Q2 integer
        q2 = result['questions'][1]
        assert q2['type'] == 'integer'
        assert q2['answer'] == '50'
        assert q2['solution'] == '10*5 equals 50'
        
        print("✅ Excel two sheets test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_two_files_questions_and_solutions():
    """Test separate Questions.xlsx and Solutions.xlsx files."""
    q_sheets = {
        'Questions': [
            ['question_no', 'subject', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'type'],
            ['1', 'Physics', 'Force unit?', 'Newton', 'Joule', 'Watt', 'Pascal', 'mcq'],
            ['2', 'Physics', 'Speed of light in m/s?', '', '', '', '', 'integer'],
        ]
    }
    
    s_sheets = {
        'Solutions': [
            ['question_no', 'solution_text', 'correct_option'],
            ['1', 'Force is measured in Newtons', 'A'],
            ['2', 'Speed of light is 3e8 m/s', '300000000'],
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        qraw = create_workbook_bytes(q_sheets)
        sraw = create_workbook_bytes(s_sheets)
        qpath = tmpdir / 'questions.xlsx'
        spath = tmpdir / 'solutions.xlsx'
        qpath.write_bytes(qraw)
        spath.write_bytes(sraw)
        
        result = await organise_pdfs(qpath, spath, 'xlsx')
        
        assert 'error' not in result
        assert len(result['questions']) == 2
        assert len(result['issues']) == 0
        
        # Verify matching by question_no
        q1 = result['questions'][0]
        assert q1['answer'] == 'A'
        assert q1['solution'] == 'Force is measured in Newtons'
        
        q2 = result['questions'][1]
        assert q2['answer'] == '300000000'
        assert q2['type'] == 'integer'
        
        print("✅ Excel two files test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_duplicate_question_id_detected():
    """Test that duplicate question IDs are flagged."""
    sheets = {
        'Questions': [
            ['question_id', 'question_no', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option', 'type', 'solution_text'],
            ['dup-1', '1', 'First question', 'A', 'B', 'C', 'D', 'A', 'mcq', 'Solution 1'],
            ['dup-1', '2', 'Second question', 'A', 'B', 'C', 'D', 'B', 'mcq', 'Solution 2'],  # Duplicate ID
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert len(result['issues']) > 0
        
        # Check for duplicate issue
        issue_messages = [i['message'].lower() for i in result['issues']]
        assert any('duplicate' in msg for msg in issue_messages)
        
        print("✅ Excel duplicate ID detection test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_missing_solution_detected():
    """Test that missing solutions are flagged."""
    sheets = {
        'Questions': [
            ['question_no', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'type'],
            ['1', 'Question 1', 'A', 'B', 'C', 'D', 'mcq'],
            ['2', 'Question 2', 'A', 'B', 'C', 'D', 'mcq'],
        ],
        'Solutions': [
            ['question_no', 'solution_text', 'correct_option'],
            ['1', 'Solution for Q1', 'A'],
            # Q2 solution missing
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert len(result['issues']) > 0
        
        # Q2 should have issue about missing solution
        issue_messages = [i['message'].lower() for i in result['issues']]
        assert any('solution' in msg for msg in issue_messages)
        
        print("✅ Excel missing solution detection test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_exact_integer_validation():
    """Test that integer answers must be exact integers (no decimals)."""
    sheets = {
        'Questions': [
            ['question_no', 'question_text', 'correct_option', 'type', 'solution_text'],
            ['1', 'What is 5*5?', '25', 'integer', 'Answer is 25'],
            ['2', 'What is pi?', '3.14', 'integer', 'Answer is 3.14'],  # Decimal not allowed
            ['3', 'What is -10?', '-10', 'integer', 'Answer is -10'],  # Negative OK
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        
        # Q1 should be OK
        q1 = result['questions'][0]
        assert q1['answer'] == '25'
        
        # Q2 should have issue (decimal)
        q2_issues = [i for i in result['issues'] if i.get('row') == 3]
        assert len(q2_issues) > 0, "Q2 should have issue for decimal in integer field"
        
        # Q3 should be OK (negative)
        q3 = result['questions'][2]
        assert q3['answer'] == '-10'
        
        print("✅ Excel exact integer validation test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_bitsat_exam_validation():
    """Test BITSAT exam only allows MCQs."""
    sheets = {
        'Questions': [
            ['question_no', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option', 'type', 'solution_text', 'exam'],
            ['1', 'MCQ question', 'A', 'B', 'C', 'D', 'A', 'mcq', 'Solution', 'BITSAT'],
            ['2', 'Integer question', '', '', '', '', '42', 'integer', 'Solution', 'BITSAT'],  # Not allowed
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert result['exam'] == 'BITSAT'
        
        # Should have issue about BITSAT not supporting integer
        issue_messages = [i['message'].lower() for i in result['issues']]
        assert any('bitsat' in msg and ('mcq' in msg or 'integer' in msg) for msg in issue_messages)
        
        print("✅ Excel BITSAT exam validation test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_formula_rejection():
    """Test that formulas are rejected."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Questions')
    
    # Add headers
    ws.append(['question_no', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'correct_option', 'type', 'solution_text'])
    
    # Add row with formula
    ws.append(['1', 'Question', 'A', 'B', 'C', 'D', 'A', 'mcq', 'Solution'])
    ws['B2'] = '=SUM(A1:A10)'  # Formula in question_text
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        out = io.BytesIO()
        wb.save(out)
        raw = out.getvalue()
        
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert len(result['issues']) > 0
        
        issue_messages = [i['message'].lower() for i in result['issues']]
        assert any('formula' in msg for msg in issue_messages)
        
        print("✅ Excel formula rejection test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.mark.asyncio
async def test_excel_unmatched_solution_row():
    """Test that unmatched solution rows are flagged."""
    sheets = {
        'Questions': [
            ['question_no', 'question_text', 'option_a', 'option_b', 'option_c', 'option_d', 'type'],
            ['1', 'Question 1', 'A', 'B', 'C', 'D', 'mcq'],
        ],
        'Solutions': [
            ['question_no', 'solution_text', 'correct_option'],
            ['1', 'Solution 1', 'A'],
            ['2', 'Solution 2', 'B'],  # No matching question
        ]
    }
    
    tmpdir = Path(tempfile.mkdtemp(prefix='excel_test_'))
    try:
        raw = create_workbook_bytes(sheets)
        qpath = tmpdir / 'questions.xlsx'
        qpath.write_bytes(raw)
        
        result = await organise_pdfs(qpath, None, 'xlsx')
        
        assert 'error' not in result
        assert len(result['issues']) > 0
        
        issue_messages = [i['message'].lower() for i in result['issues']]
        assert any('unmatched' in msg for msg in issue_messages)
        
        print("✅ Excel unmatched solution row detection test passed")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == '__main__':
    pytest.main([__file__, '-v', '-s'])
